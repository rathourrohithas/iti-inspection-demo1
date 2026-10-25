"""Excel and CSV exports. Field names and numbering follow the approved proforma."""
import csv
import io

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .proforma import CORRECTIONS, J2_ROWS, PROFORMA
from .schema import build_index, j2_summary, pct
from .services import team_rows
from .tables import build, fmt

HEAD_FILL = PatternFill("solid", fgColor="E8EDF4")
TITLE_FONT = Font(bold=True, size=12, color="1F3A5F")
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Side(style="thin", color="9AA9BD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _cover_rows(insp):
    team = team_rows(insp)
    leader = next((t for t in team if t["role"] == "LEADER"), None)
    return [
        ("Name of ITI", insp.iti.name), ("MIS code", insp.iti.mis_code), ("District", insp.iti.district),
        ("Ownership", insp.iti.get_ownership_display()), ("Inspection team", insp.team_name),
        ("Team leader", leader["name"] if leader else ""),
        ("Team members", ", ".join(t["name"] for t in team if t["role"] != "LEADER")),
        ("Scheduled visit", insp.scheduled_date.strftime("%d-%m-%Y")),
        ("Date of inspection", fmt((insp.data or {}).get("SA.A7.date"), "date")),
        ("Submission deadline", insp.deadline.strftime("%d-%m-%Y")),
        ("Status", insp.get_status_display()),
        ("Submitted on", timezone.localtime(insp.submitted_at).strftime("%d-%m-%Y %H:%M") if insp.submitted_at else ""),
        ("Submitted by", (insp.submitted_by.get_full_name() or insp.submitted_by.username) if insp.submitted_by else ""),
        ("Proforma version", insp.proforma_version or PROFORMA["version"]),
        ("Generated", timezone.localtime().strftime("%d-%m-%Y %H:%M")),
    ]


def inspection_workbook(insp):
    """One sheet per Part, laid out like the proforma tables, plus cover and attachments."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Cover"
    ws.append([PROFORMA["title"]])
    ws["A1"].font = Font(bold=True, size=14, color="1F3A5F")
    ws.append(["Directorate of Technical Education, Vocational & Industrial Training, Himachal Pradesh"])
    ws.append([])
    for k, v in _cover_rows(insp):
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = BOLD
    ws.append([])
    ws.append(["Notes on the digital proforma"])
    ws.cell(ws.max_row, 1).font = BOLD
    for c in CORRECTIONS:
        ws.append(["", c])
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 90

    atts = {}
    att_list = list(insp.attachments.filter(deleted=False).select_related("uploaded_by"))
    for a in att_list:
        atts.setdefault(a.unit_key, []).append(a)

    used = set()
    for sec_title, part_title, els in build(insp, atts):
        sec_prefix = "A" if sec_title.startswith("Section A") else "B"
        code = part_title.split("—")[0].replace("Part", "").strip().rstrip(".").split(".")[0].strip()
        name = f"{sec_prefix}-{code} " + part_title.split("—")[-1].split(".")[-1].strip()
        name = "".join(ch for ch in name if ch not in "[]:*?/\\")[:31]
        while name in used:
            name = name[:29] + "_2"
        used.add(name)
        ws = wb.create_sheet(name)
        ws.append([sec_title])
        ws.cell(1, 1).font = Font(bold=True, color="E8A33D")
        ws.append([part_title])
        ws.cell(2, 1).font = TITLE_FONT
        maxcols = 1
        widths = {}
        for e in els:
            if e[0] in ("note", "para"):
                ws.append([e[1]])
                ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=False, vertical="top")
                ws.cell(ws.max_row, 1).font = Font(italic=e[0] == "note", color="3D4A5C")
                continue
            tb = e[1]
            ws.append([])
            if tb.title:
                ws.append([tb.title])
                ws.cell(ws.max_row, 1).font = Font(bold=True, color="2C5282")
            if any(tb.headers):
                ws.append(tb.headers)
                for i in range(1, len(tb.headers) + 1):
                    c = ws.cell(ws.max_row, i)
                    c.font = BOLD
                    c.fill = HEAD_FILL
                    c.alignment = WRAP
                    c.border = BORDER
            for r in tb.rows:
                ws.append(r)
                for i in range(1, len(r) + 1):
                    c = ws.cell(ws.max_row, i)
                    c.alignment = WRAP
                    c.border = BORDER
            maxcols = max(maxcols, len(tb.headers))
            rel = tb.widths or [100 / len(tb.headers)] * len(tb.headers)
            for i, w in enumerate(rel, 1):
                widths[i] = max(widths.get(i, 0), min(70, max(9, w * 1.25)))
        for i in range(1, maxcols + 1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(i, 14)
        ws.sheet_view.zoomScale = 90

    ws = wb.create_sheet("Attachments")
    ws.append(["Item", "File name", "Caption", "Size (KB)", "Uploaded by", "Uploaded on", "Link (requires login)"])
    for i in range(1, 8):
        ws.cell(1, i).font = BOLD
        ws.cell(1, i).fill = HEAD_FILL
    ub = {u.key: u for u in build_index()[0]}
    for a in att_list:
        u = ub.get(a.unit_key)
        ws.append([f"{u.sno} {u.label}" if u else a.unit_key, a.original_name, a.caption, round(a.size / 1024),
                   (a.uploaded_by.get_full_name() or a.uploaded_by.username) if a.uploaded_by else "",
                   timezone.localtime(a.uploaded_at).strftime("%d-%m-%Y %H:%M"), f"/attachment/{a.id}/"])
    for col, w in zip("ABCDEFG", (50, 30, 30, 10, 22, 18, 24)):
        ws.column_dimensions[col].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------- long format
LONG_HEADERS = ["Inspection ID", "Name of ITI", "MIS code", "District", "Inspection team", "Team leader",
                "Scheduled visit", "Date of inspection", "Status", "Submitted on",
                "Section", "Part", "S.No", "Item", "Field", "Row", "Value"]


def _input_label(spec_key, unit, col_label=None):
    suffix = spec_key[len(unit.key):].lstrip(".")
    names = {"obs": "Observed", "cpn": "C/P/N", "measured": "Measured area / count", "maint": "Maintained (Y/N)",
             "uptodate": "Up to date (Y/N)", "remarks": "Remarks", "avail": "Available (Yes/No)", "updated": "Updated up to",
             "deficiency": "Deficiency observed", "action": "Action required", "found": "Found (Y/N)",
             "evidence": "Evidence", "members": "Declaration", "ack_received": "Copy received by Principal on",
             "atr_due": "Action Taken Report will be submitted by", "sanctioned": "Sanctioned / required",
             "in_position": "In position", "vacant": "Vacant", "sanctioned_kw": "Sanctioned (kW)",
             "required_kw": "Required (kW)", "required": "Required", "working": "Working",
             "rs_per_trainee": "Expenditure per trainee per month (Rs)", "last_report": "Last report dated"}
    if not suffix:
        return col_label or "Response"
    if suffix.startswith("hdr."):
        return "Heading: " + suffix[4:]
    return names.get(suffix, suffix)


def _column_labels():
    """Column labels for list-type inputs, from the schema."""
    labels = {}
    for sec in PROFORMA["sections"]:
        for part in sec["parts"]:
            for b in part["blocks"]:
                if b["type"] == "grid":
                    labels[f"{sec['id']}.{b['id']}"] = {c["k"]: c["label"] for c in b["columns"]}
                elif b["type"] == "particulars":
                    for r in b["rows"]:
                        for f in r["fields"]:
                            if f["t"] == "rows":
                                labels[f"{sec['id']}.{r['sno']}"] = {c["k"]: c["label"] for c in f["columns"]}
                elif b["type"] == "staffing":
                    d = {"trade": "Trade"}
                    d.update({c["k"]: c["label"] for c in b["columns"]})
                    labels[f"{sec['id']}.{b['id']}.trades"] = d
                elif b["type"] == "area":
                    for r in b["rows"]:
                        if r.get("repeat"):
                            labels[f"{sec['id']}.{r['sno']}"] = {"trade_unit": "Trade / unit", "measured": "Measured area / count", "cpn": "C/P/N"}
                elif b["type"] == "declaration":
                    labels[f"{sec['id']}.{b['id']}.members"] = {"name": "Name", "designation": "Designation and office",
                                                                "declared": "Declared", "declared_at": "Declared on",
                                                                "declared_by": "Declared by", "date": "Date"}
    return labels


def _iter_long_rows(insps):
    units, inputs, parts = build_index()
    part_titles = {p["key"]: (p["section_title"], p["title"]) for p in parts}
    col_labels = _column_labels()
    for insp in insps:
        team = team_rows(insp)
        leader = next((t["name"] for t in team if t["role"] == "LEADER"), "")
        base = [insp.id, insp.iti.name, insp.iti.mis_code, insp.iti.district, insp.team_name, leader,
                insp.scheduled_date.strftime("%d-%m-%Y"), fmt((insp.data or {}).get("SA.A7.date"), "date"),
                insp.get_status_display(),
                timezone.localtime(insp.submitted_at).strftime("%d-%m-%Y %H:%M") if insp.submitted_at else ""]
        data = insp.data or {}
        for u in units:
            sec_t, part_t = part_titles[u.part_key]
            if u.part_key.startswith("SB.F") and not insp.iti.is_govt:
                continue
            for key in u.inputs:
                v = data.get(key)
                spec = inputs[key]
                if isinstance(v, list):
                    cl = col_labels.get(key, {})
                    n = 0
                    for row in v:
                        if not isinstance(row, dict) or not any(x not in (None, "", False) for k, x in row.items() if k != "user_id"):
                            continue
                        n += 1
                        if key == "SA.G1":
                            for y in ("y1", "y2", "y3"):
                                p = pct(row.get(f"{y}_pass"), row.get(f"{y}_app"))
                                if p != "":
                                    row = dict(row, **{f"{y}_pct": p})
                        for ck, cv in row.items():
                            if ck == "user_id" or cv in (None, ""):
                                continue
                            yield base + [sec_t, part_t, u.sno, u.label, cl.get(ck, ck), n,
                                          fmt(cv, spec["columns"].get(ck, {}).get("t"))]
                elif v not in (None, ""):
                    yield base + [sec_t, part_t, u.sno, u.label, _input_label(key, u), "", fmt(v, spec["t"])]
        for r in j2_summary(data):
            for k in ("checked", "C", "P", "N", "NA"):
                yield base + ["Section A – Technical Evaluation", "Part J — Findings, recommendation and certification",
                              "J2", "J2. Summary of compliance by part", f"{r['label']} — {'Items checked' if k == 'checked' else k}", "", r[k]]


def long_export_csv(insps):
    buf = io.StringIO()
    buf.write("﻿")  # Excel-friendly UTF-8
    w = csv.writer(buf)
    w.writerow(LONG_HEADERS)
    for row in _iter_long_rows(insps):
        w.writerow(row)
    return buf.getvalue().encode("utf-8")


def _cpn_matrix_columns():
    cols = []
    for u in build_index()[0]:
        if u.section_id != "SA":
            continue
        if u.kind in ("checklist",) or (u.kind == "area"):
            cols.append((u, "cpn"))
        elif u.kind == "register":
            cols.append((u, "maint"))
            cols.append((u, "uptodate"))
        elif u.kind == "j1":
            cols.append((u, "found"))
    return cols


def long_export_xlsx(insps):
    insps = list(insps)
    wb = Workbook(write_only=False)
    ws = wb.active
    ws.title = "All responses"
    ws.append(LONG_HEADERS)
    for i in range(1, len(LONG_HEADERS) + 1):
        ws.cell(1, i).font = BOLD
        ws.cell(1, i).fill = HEAD_FILL
    for row in _iter_long_rows(insps):
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(LONG_HEADERS))}{ws.max_row}"
    for i, w in enumerate((8, 34, 14, 14, 20, 20, 12, 12, 14, 16, 26, 34, 8, 50, 26, 6, 40), 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # One row per inspection with each compliance response.
    ws2 = wb.create_sheet("Compliance matrix")
    cols = _cpn_matrix_columns()
    head = ["Inspection ID", "Name of ITI", "MIS code", "District", "Status", "Date of inspection"]
    for u, f in cols:
        if u.kind == "register":
            head.append(f"{u.sno} {'Maintained' if f == 'maint' else 'Up to date'}")
        elif u.kind == "j1":
            head.append(f"J1 {u.label.split(')')[0]})")
        else:
            head.append(u.sno)
    j2_labels = [r[1] for r in J2_ROWS] + ["Total"]
    for lab in j2_labels:
        for k in ("C", "P", "N", "NA"):
            head.append(f"J2 {lab.split(' —')[0]} {k}")
    head.append("J5 Overall assessment")
    ws2.append(head)
    for i in range(1, len(head) + 1):
        ws2.cell(1, i).font = BOLD
        ws2.cell(1, i).fill = HEAD_FILL
        ws2.cell(1, i).alignment = Alignment(wrap_text=True, vertical="top")
    for insp in insps:
        d = insp.data or {}
        row = [insp.id, insp.iti.name, insp.iti.mis_code, insp.iti.district, insp.get_status_display(),
               fmt(d.get("SA.A7.date"), "date")]
        for u, f in cols:
            if u.kind == "area" and isinstance(d.get(u.key), list):
                row.append(", ".join(f"{fmt(x.get('trade_unit'))}: {fmt(x.get('cpn'))}" for x in d[u.key] if isinstance(x, dict) and x.get("cpn")))
            else:
                row.append(fmt(d.get(f"{u.key}.{f}")))
        for r in j2_summary(d):
            for k in ("C", "P", "N", "NA"):
                row.append(r[k])
        j5 = d.get("SA.J5") or ""
        row.append(j5.split(" — ")[0].split(":")[0] if j5 else "")
        ws2.append(row)
    ws2.freeze_panes = "G2"
    ws2.row_dimensions[1].height = 42

    ws3 = wb.create_sheet("Notes")
    ws3.append(["Exported", timezone.localtime().strftime("%d-%m-%Y %H:%M")])
    ws3.append(["Records", len(insps)])
    ws3.append([])
    ws3.append(["'All responses' lists every answered field with the proforma's own Section, Part, S.No and wording."])
    ws3.append(["'Compliance matrix' has one row per inspection with each C/P/N, Y/N and J2 count, for comparison across ITIs."])
    ws3.append([])
    ws3.append(["Notes on the digital proforma"])
    for c in CORRECTIONS:
        ws3.append([c])
    ws3.column_dimensions["A"].width = 120
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

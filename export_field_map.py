"""Write docs/FIELD_MAP.xlsx: every proforma item mapped to its digital input(s)."""
from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from inspections.proforma import CORRECTIONS, PROFORMA
from inspections.schema import build_index

TYPE_NAMES = {"text": "Text", "textarea": "Text (multi-line)", "number": "Number", "date": "Date", "time": "Time",
              "email": "E-mail", "tel": "Phone", "lat": "Latitude", "lng": "Longitude", "radio": "Single choice",
              "select": "Drop-down", "grid": "Table (add rows)", "rows": "Table (add rows)", "area_rows": "Table (add rows)",
              "trades": "Table (add rows)", "members": "Declaration table", "bool": "Tick box", "int": "Number"}


class Command(BaseCommand):
    help = "Export the field map of the digital proforma to docs/FIELD_MAP.xlsx"

    def handle(self, *a, **k):
        units, inputs, parts = build_index()
        ptitle = {p["key"]: (p["section_title"], p["title"]) for p in parts}
        norms = {}
        for sec in PROFORMA["sections"]:
            for part in sec["parts"]:
                for b in part["blocks"]:
                    for r in b.get("rows", []):
                        if isinstance(r, dict) and r.get("sno"):
                            norms[f"{sec['id']}.{r['sno']}"] = r.get("norm", "")
                        if isinstance(r, dict) and r.get("k") and b["type"] == "j1":
                            norms[f"{sec['id']}.J1.{r['k']}"] = "Linked items: " + r["linked"]
                    for p in b.get("posts", []):
                        norms[f"{sec['id']}.{b['id']}.{p['k']}"] = p["norm"]
        wb = Workbook()
        ws = wb.active
        ws.title = "Field map"
        head = ["Section", "Part", "S.No", "Item (as in proforma)", "Norm / Ref", "Field key", "Response captured",
                "Input type", "Options / columns", "Photo / document upload"]
        ws.append(head)
        for i in range(1, len(head) + 1):
            ws.cell(1, i).font = Font(bold=True)
            ws.cell(1, i).fill = PatternFill("solid", fgColor="E8EDF4")
        from inspections.exports import _input_label
        for u in units:
            s, p = ptitle[u.part_key]
            for key in u.inputs:
                spec = inputs[key]
                opts = ", ".join(spec.get("options") or []) if spec.get("options") else ""
                if spec.get("columns"):
                    opts = "Columns: " + ", ".join(spec["columns"].keys())
                ws.append([s, p, u.sno, u.label, norms.get(u.key, ""), key, _input_label(key, u), TYPE_NAMES.get(spec["t"], spec["t"]),
                           opts, "Optional" if key == u.inputs[0] else ""])
        for col, w in zip("ABCDEFGHIJ", (22, 34, 8, 60, 40, 22, 24, 16, 30, 14)):
            ws.column_dimensions[col].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        ws2 = wb.create_sheet("Corrections")
        ws2.append(["Corrections applied while digitising (approved)"])
        ws2.cell(1, 1).font = Font(bold=True)
        for c in CORRECTIONS:
            ws2.append([c])
        ws2.column_dimensions["A"].width = 130
        ws3 = wb.create_sheet("Summary")
        ws3.append(["Answerable items", len(units)])
        ws3.append(["Input fields", len(inputs)])
        ws3.append(["Parts", len(parts)])
        out = Path(settings.BASE_DIR) / "docs" / "FIELD_MAP.xlsx"
        wb.save(out)
        self.stdout.write(f"Wrote {out} ({len(units)} items, {len(inputs)} inputs)")

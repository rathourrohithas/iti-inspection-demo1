import datetime
import io
import json
import re

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .models import ITI, Attachment, AuditLog, Inspection, PortalSettings, Profile, TeamMember
from .proforma import PROFORMA
from .schema import build_index, j2_summary
from .services import seed_prefill

User = get_user_model()
PW = "Strong#Pass2026x"


def mkuser(username, role=Profile.INSPECTOR):
    u = User.objects.create_user(username=username, password=PW, first_name=username.title())
    u.profile.role = role
    u.profile.save()
    return u


class FidelityTests(TestCase):
    """The digital form must contain every item of the approved proforma."""

    def test_every_numbered_item_present(self):
        units, inputs, parts = build_index()
        snos = {u.sno for u in units}
        expected = (
            [f"A{i}" for i in range(1, 15)] + [f"B2.{i}" for i in range(1, 9)] + [f"C1.{i}" for i in range(1, 14)]
            + [f"C2.{i}" for i in range(1, 17)] + [f"C3.{i}" for i in range(1, 7)] + [f"C4.{i}" for i in range(1, 10)]
            + [f"D2.{i}" for i in range(1, 16)] + [f"E3.{i}" for i in range(1, 7)] + [f"F3.{i}" for i in range(1, 10)]
            + [f"F4.{i}" for i in range(1, 14)] + [f"F5.{i}" for i in range(1, 6)] + [f"G2.{i}" for i in range(1, 7)]
            + [f"H1.{i}" for i in range(1, 10)] + [f"H2.{i}" for i in range(1, 9)] + [f"H3.{i}" for i in range(1, 7)]
            + [f"H4.{i}" for i in range(1, 6)] + [f"H5.{i}" for i in range(1, 3)] + [f"H6.{i}" for i in range(1, 6)]
            + [f"H7.{i}" for i in range(1, 15)] + [f"I1.{i}" for i in range(1, 17)] + [f"I2.{i}" for i in range(1, 4)]
            + ["B1", "D1", "DTC", "E2", "F1", "F2", "FCO", "G1", "J3", "J4", "J5", "J6", "J7", "J8"]
        )
        missing = [s for s in expected if s not in snos]
        self.assertEqual(missing, [])
        # Section B: running numbers 1-74
        sb = sorted(int(u.sno.split()[-1]) for u in units if u.section_id == "SB")
        self.assertEqual(sb, list(range(1, 75)))
        # J1: ten grounds; E1: eleven posts
        self.assertEqual(len([u for u in units if u.kind == "j1"]), 10)
        self.assertEqual(len([u for u in units if u.kind == "staffing"]), 11)
        self.assertEqual(len(parts), 17)

    def test_cpn_options_include_na(self):
        self.assertEqual(PROFORMA["cpn_options"], ["C", "P", "N", "NA"])


class WorkflowTests(TestCase):
    def setUp(self):
        self.admin = mkuser("admin1", Profile.ADMIN)
        self.viewer = mkuser("viewer1", Profile.VIEWER)
        self.lead = mkuser("lead1")
        self.mem = mkuser("mem1")
        self.outsider = mkuser("other1")
        self.iti = ITI.objects.create(name="Govt. ITI Test", mis_code="T-1", district="Mandi", address="Addr")
        self.insp = Inspection.objects.create(iti=self.iti, team_name="T1", scheduled_date=timezone.localdate(),
                                              deadline=timezone.localdate() + datetime.timedelta(days=3), created_by=self.admin)
        TeamMember.objects.create(inspection=self.insp, user=self.lead, role=TeamMember.LEADER)
        self.tm = TeamMember.objects.create(inspection=self.insp, user=self.mem, role=TeamMember.MEMBER, sections=["SA.C"])
        seed_prefill(self.insp, refresh_team=True)
        self.insp.save()

    def login(self, u):
        self.client.force_login(u)

    def save(self, changes, user=None):
        if user:
            self.login(user)
        return self.client.post(reverse("api_save", args=[self.insp.id]), json.dumps({"changes": changes}),
                                content_type="application/json")

    def test_prefill(self):
        self.assertEqual(self.insp.data["SA.A1"], "Govt. ITI Test")
        self.assertEqual(len(self.insp.data["SA.J8.members"]), 2)

    def test_member_section_restriction(self):
        r = self.save({"SA.C1.1.cpn": "C", "SA.B2.1.cpn": "C"}, self.mem).json()
        self.assertIn("SA.C1.1.cpn", r["saved"])
        self.assertIn("SA.B2.1.cpn", r["errors"])
        self.insp.refresh_from_db()
        self.assertEqual(self.insp.status, Inspection.IN_PROGRESS)
        self.assertNotIn("SA.B2.1.cpn", self.insp.data)

    def test_outsider_cannot_view_or_save(self):
        self.login(self.outsider)
        self.assertEqual(self.client.get(reverse("record_view", args=[self.insp.id])).status_code, 403)
        self.assertEqual(self.save({"SA.A4": "Yes"}).status_code, 403)
        self.assertEqual(self.client.get(reverse("dashboard")).status_code, 403)

    def test_validation(self):
        r = self.save({"SA.A7.date": "31-02-2026", "SA.B1": [{"seats": "abc"}], "SA.A4": "Maybe", "SA.A5.email": "x"}, self.lead).json()
        self.assertEqual(set(r["errors"]), {"SA.A7.date", "SA.B1", "SA.A4", "SA.A5.email"})

    def test_g1_pass_percentage_calculated(self):
        self.save({"SA.G1": [{"trade": "Fitter", "y1_app": 40, "y1_pass": 30}]}, self.lead)
        self.insp.refresh_from_db()
        self.assertEqual(self.insp.data["SA.G1"][0]["y1_pct"], 75.0)

    def test_j2_counts(self):
        self.save({"SA.B2.1.cpn": "C", "SA.B2.2.cpn": "P", "SA.C1.1.cpn": "NA", "SA.I1.1.maint": "Y", "SA.I1.1.uptodate": "N",
                   "SA.I2.3.cpn": "N", "SA.C2.1": [{"trade_unit": "Fitter U1", "cpn": "C"}, {"trade_unit": "Fitter U2", "cpn": "N"}]}, self.lead)
        self.insp.refresh_from_db()
        j2 = {r["code"]: r for r in j2_summary(self.insp.data)}
        self.assertEqual((j2["B"]["C"], j2["B"]["P"]), (1, 1))
        self.assertEqual((j2["C"]["NA"], j2["C"]["C"], j2["C"]["N"], j2["C"]["checked"]), (1, 1, 1, 3))
        self.assertEqual((j2["I"]["P"], j2["I"]["N"]), (1, 1))
        self.assertEqual(j2["Total"]["checked"], 7)

    def test_declaration_only_own_row(self):
        members = self.insp.data["SA.J8.members"]
        forged = [dict(m, declared=True) for m in members]
        self.save({"SA.J8.members": forged}, self.lead)
        self.insp.refresh_from_db()
        by = {m["user_id"]: m for m in self.insp.data["SA.J8.members"]}
        self.assertTrue(by[self.lead.id]["declared"])
        self.assertTrue(by[self.lead.id]["declared_at"])
        self.assertFalse(by[self.mem.id]["declared"])

    def test_mandatory_blocks_submission_and_member_cannot_submit(self):
        ps = PortalSettings.get()
        ps.mandatory_units = ["SA.A4"]
        ps.save()
        self.login(self.mem)
        self.assertEqual(self.client.post(reverse("inspection_submit", args=[self.insp.id]), {"confirm": "yes"}).status_code, 403)
        self.login(self.lead)
        self.client.post(reverse("inspection_submit", args=[self.insp.id]), {"confirm": "yes"})
        self.insp.refresh_from_db()
        self.assertNotEqual(self.insp.status, Inspection.SUBMITTED)
        self.save({"SA.A4": "No"})
        self.client.post(reverse("inspection_submit", args=[self.insp.id]), {"confirm": "yes"})
        self.insp.refresh_from_db()
        self.assertEqual(self.insp.status, Inspection.SUBMITTED)

    def test_submit_locks_and_prevents_duplicates_and_return_reopens(self):
        self.login(self.lead)
        self.save({"SA.A4": "Yes"})
        self.client.post(reverse("inspection_submit", args=[self.insp.id]), {"confirm": "yes"})
        self.client.post(reverse("inspection_submit", args=[self.insp.id]), {"confirm": "yes"})
        self.assertEqual(AuditLog.objects.filter(inspection=self.insp, action="SUBMITTED").count(), 1)
        r = self.save({"SA.A4": "No"})
        self.assertEqual(r.status_code, 409)
        self.login(self.lead)
        self.assertRedirects(self.client.get(reverse("inspection_form", args=[self.insp.id])), reverse("record_view", args=[self.insp.id]))
        # Only the administrator can return
        self.client.post(reverse("inspection_return", args=[self.insp.id]), {"reason": "x"})
        self.insp.refresh_from_db()
        self.assertEqual(self.insp.status, Inspection.SUBMITTED)
        self.login(self.admin)
        self.client.post(reverse("inspection_return", args=[self.insp.id]), {"reason": "Fill C2"})
        self.insp.refresh_from_db()
        self.assertEqual(self.insp.status, Inspection.RETURNED)
        self.assertEqual(self.save({"SA.A4": "No"}, self.lead).status_code, 200)

    def test_audit_trail_records_field_changes(self):
        self.save({"SA.A4": "Yes"}, self.lead)
        self.save({"SA.A4": "No"})
        log = AuditLog.objects.filter(inspection=self.insp, action="FIELD_CHANGED", field_key="SA.A4").order_by("timestamp")
        self.assertEqual([(l.old_value, l.new_value) for l in log], [("", "Yes"), ("Yes", "No")])
        self.assertEqual(log[0].user, self.lead)

    @override_settings(MEDIA_ROOT="/tmp/iti_test_media")
    def test_upload_rules(self):
        self.login(self.mem)
        jpg = b"\xff\xd8\xff\xe0" + b"0" * 100
        url = reverse("api_upload", args=[self.insp.id])
        ok = self.client.post(url, {"unit_key": "SA.C1.1", "file": SimpleUploadedFile("a.jpg", jpg, "image/jpeg")}).json()
        self.assertTrue(ok["ok"])
        bad_section = self.client.post(url, {"unit_key": "SA.B2.1", "file": SimpleUploadedFile("a.jpg", jpg, "image/jpeg")})
        self.assertEqual(bad_section.status_code, 403)
        fake = self.client.post(url, {"unit_key": "SA.C1.1", "file": SimpleUploadedFile("a.pdf", b"MZ....", "application/pdf")})
        self.assertEqual(fake.status_code, 400)
        exe = self.client.post(url, {"unit_key": "SA.C1.1", "file": SimpleUploadedFile("a.exe", b"MZ", "application/x")})
        self.assertEqual(exe.status_code, 400)
        att = Attachment.objects.get()
        self.login(self.outsider)
        self.assertEqual(self.client.get(reverse("attachment_file", args=[att.id])).status_code, 403)
        self.login(self.viewer)
        self.assertEqual(self.client.get(reverse("attachment_file", args=[att.id])).status_code, 200)

    def test_reports_and_exports(self):
        self.save({"SA.A4": "Yes", "SA.B2.1.cpn": "C", "SA.B2.1.obs": "Order seen", "SA.J5": PROFORMA["sections"][0]["parts"][9]["blocks"][4]["options"][0]}, self.lead)
        self.login(self.viewer)
        pdf = self.client.get(reverse("inspection_pdf", args=[self.insp.id]))
        self.assertEqual(pdf["Content-Type"], "application/pdf")
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        x = self.client.get(reverse("inspection_xlsx", args=[self.insp.id]))
        wb = load_workbook(io.BytesIO(x.content))
        self.assertIn("Cover", wb.sheetnames)
        self.assertEqual(len(wb.sheetnames), 1 + 17 + 1)
        csvr = self.client.get(reverse("records") + "?export=csv").content.decode("utf-8-sig")
        self.assertIn("B2.1", csvr)
        self.assertIn("Order seen", csvr)
        xl = self.client.get(reverse("records") + "?export=xlsx")
        wb = load_workbook(io.BytesIO(xl.content))
        self.assertEqual(wb.sheetnames, ["All responses", "Compliance matrix", "Notes"])

    def test_dashboard_and_filters(self):
        self.login(self.viewer)
        r = self.client.get(reverse("dashboard") + "?district=Mandi&status=NOT_STARTED")
        self.assertContains(r, "Govt. ITI Test")
        r = self.client.get(reverse("dashboard") + "?status=SUBMITTED")
        self.assertNotContains(r, "Govt. ITI Test</td>")
        # Viewer is read-only for administration
        self.assertEqual(self.client.get(reverse("users_list")).status_code, 403)

    def test_inspector_sees_only_own(self):
        self.login(self.outsider)
        r = self.client.get(reverse("my_inspections"))
        self.assertNotContains(r, "Govt. ITI Test")
        self.login(self.mem)
        self.assertContains(self.client.get(reverse("my_inspections")), "Govt. ITI Test")
        self.assertEqual(self.client.get(reverse("inspection_form", args=[self.insp.id])).status_code, 200)


class AdminTests(TestCase):
    def setUp(self):
        self.admin = mkuser("admin1", Profile.ADMIN)
        self.client.force_login(self.admin)

    def test_create_user_forces_password_change(self):
        r = self.client.post(reverse("user_create"), {"username": "new.officer", "full_name": "New Officer", "role": Profile.INSPECTOR,
                                                      "is_active": "on", "password": "Temp#Pass2026x"})
        self.assertEqual(r.status_code, 302)
        u = User.objects.get(username="new.officer")
        self.client.force_login(u)
        self.assertRedirects(self.client.get(reverse("my_inspections")), reverse("password_change"))

    def test_import_itis_and_schedule(self):
        mkuser("lead.x")
        mkuser("mem.x")
        csv_itis = "mis_code,name,ownership,district,block\nX-1,Govt ITI X,GOVT,Kullu,Banjar\nX-2,Pvt ITI Y,PRIVATE,Solan,Baddi\n,missing,,,\n"
        r = self.client.post(reverse("iti_import"), {"file": SimpleUploadedFile("itis.csv", csv_itis.encode())})
        self.assertContains(r, "Created: <b>2</b>")
        self.assertEqual(ITI.objects.get(mis_code="X-2").ownership, ITI.PRIVATE)
        csv_s = ("mis_code,team_name,scheduled_date,deadline,leader_username,member_usernames\n"
                 "X-1,Team K,10-11-2026,2026-11-15,lead.x,mem.x\nX-9,Team Z,2026-11-10,2026-11-15,lead.x,\n")
        r = self.client.post(reverse("assignment_import"), {"file": SimpleUploadedFile("s.csv", csv_s.encode())})
        self.assertContains(r, "Created: <b>1</b>")
        self.assertContains(r, "no ITI with MIS code")
        insp = Inspection.objects.get(iti__mis_code="X-1")
        self.assertEqual(insp.team.count(), 2)
        self.assertEqual(insp.leader().username, "lead.x")

    def test_private_iti_hides_swf(self):
        iti = ITI.objects.create(name="Pvt", mis_code="P-1", district="Solan", ownership=ITI.PRIVATE)
        insp = Inspection.objects.create(iti=iti, team_name="t", scheduled_date=timezone.localdate(), deadline=timezone.localdate())
        TeamMember.objects.create(inspection=insp, user=self.admin, role=TeamMember.LEADER)
        r = self.client.post(reverse("api_save", args=[insp.id]), json.dumps({"changes": {"SB.61.avail": "Yes"}}), content_type="application/json")
        self.assertIn("SB.61.avail", r.json()["errors"])

    def test_login_lockout(self):
        self.client.logout()
        for _ in range(6):
            r = self.client.post(reverse("login"), {"username": "admin1", "password": "wrong"})
        r = self.client.post(reverse("login"), {"username": "admin1", "password": PW})
        self.assertContains(r, "Too many failed attempts")

    def test_settings_mandatory(self):
        r = self.client.post(reverse("portal_settings"), {"max_upload_mb": 5, "allowed_types": ["jpg", "pdf"], "mandatory": ["SA.A1", "SA.B2.1", "bogus"]})
        self.assertEqual(r.status_code, 302)
        ps = PortalSettings.get()
        self.assertEqual(ps.mandatory_units, ["SA.A1", "SA.B2.1"])
        self.assertEqual(ps.allowed_ext_list(), ["jpg", "pdf"])

import csv
import io

from django.contrib import messages
from django.contrib.auth import get_user_model
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Count, Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST

from . import audit
from .access import admin_required, directorate_required
from .forms import AssignmentForm, ImportForm, ITIForm, SettingsForm, UserForm
from .models import ITI, AuditLog, Inspection, PortalSettings, Profile, TeamMember
from .schema import build_index, parts_list
from .services import parse_date, seed_prefill

User = get_user_model()


# ------------------------------------------------------------------ filtering
def filter_inspections(request, qs=None):
    qs = (qs if qs is not None else Inspection.objects.all()).select_related("iti").prefetch_related(
        "team__user")
    f = {k: (request.GET.get(k) or "").strip() for k in ("district", "iti", "team", "status", "date_from", "date_to")}
    if f["district"]:
        qs = qs.filter(iti__district=f["district"])
    if f["iti"]:
        qs = qs.filter(Q(iti__name__icontains=f["iti"]) | Q(iti__mis_code__icontains=f["iti"]))
    if f["team"]:
        qs = qs.filter(Q(team_name__icontains=f["team"]) | Q(team__user__first_name__icontains=f["team"])
                       | Q(team__user__last_name__icontains=f["team"])).distinct()
    today = timezone.localdate()
    if f["status"] == "OVERDUE":
        qs = qs.exclude(status=Inspection.SUBMITTED).filter(deadline__lt=today)
    elif f["status"]:
        qs = qs.filter(status=f["status"])
    for key, lookup in (("date_from", "scheduled_date__gte"), ("date_to", "scheduled_date__lte")):
        if f[key]:
            try:
                qs = qs.filter(**{lookup: parse_date(f[key])})
            except ValueError:
                pass
    return qs, f


def _filter_choices():
    return {
        "districts": ITI.objects.values_list("district", flat=True).distinct().order_by("district"),
        "statuses": Inspection.STATUS_CHOICES + [("OVERDUE", "Pending beyond deadline")],
    }


# ------------------------------------------------------------------ dashboard
@login_required
@directorate_required
def dashboard(request):
    qs, f = filter_inspections(request)
    today = timezone.localdate()
    items = list(qs.order_by("deadline", "iti__name"))
    stats = {
        "total": len(items),
        "not_started": sum(1 for i in items if i.status == Inspection.NOT_STARTED),
        "in_progress": sum(1 for i in items if i.status in (Inspection.IN_PROGRESS, Inspection.RETURNED)),
        "returned": sum(1 for i in items if i.status == Inspection.RETURNED),
        "submitted": sum(1 for i in items if i.status == Inspection.SUBMITTED),
        "overdue": sum(1 for i in items if i.status != Inspection.SUBMITTED and i.deadline < today),
    }
    by_district = {}
    for i in items:
        d = by_district.setdefault(i.iti.district, {"district": i.iti.district, "total": 0, "NOT_STARTED": 0,
                                                     "IN_PROGRESS": 0, "RETURNED": 0, "SUBMITTED": 0, "overdue": 0})
        d["total"] += 1
        d[i.status] += 1
        if i.status != Inspection.SUBMITTED and i.deadline < today:
            d["overdue"] += 1
    by_team = {}
    for i in items:
        leader = next((t.user for t in i.team.all() if t.role == TeamMember.LEADER), None)
        key = (i.team_name, leader.id if leader else None)
        t = by_team.setdefault(key, {"team": i.team_name, "leader": (leader.get_full_name() or leader.username) if leader else "—",
                                     "assigned": 0, "submitted": 0, "progress_sum": 0, "overdue": 0})
        t["assigned"] += 1
        t["progress_sum"] += 100 if i.status == Inspection.SUBMITTED else i.progress_pct
        if i.status == Inspection.SUBMITTED:
            t["submitted"] += 1
        elif i.deadline < today:
            t["overdue"] += 1
    teams = []
    for t in by_team.values():
        t["avg"] = round(t["progress_sum"] / t["assigned"]) if t["assigned"] else 0
        teams.append(t)
    teams.sort(key=lambda t: (-t["overdue"], t["avg"]))
    for i in items:
        i.leader_obj = next((t.user for t in i.team.all() if t.role == TeamMember.LEADER), None)
    return render(request, "inspections/dashboard.html", {
        "stats": stats, "district_rows": sorted(by_district.values(), key=lambda d: d["district"]),
        "teams": teams, "items": items, "f": f, "today": today, **_filter_choices(),
    })


# ------------------------------------------------------------------ records
@login_required
@directorate_required
def records(request):
    qs, f = filter_inspections(request)
    qs = qs.order_by("-scheduled_date", "iti__name")
    fmt = request.GET.get("export")
    if fmt in ("csv", "xlsx"):
        from .exports import long_export_csv, long_export_xlsx
        ids = list(qs.values_list("id", flat=True))
        insps = Inspection.objects.filter(id__in=ids).select_related("iti").prefetch_related("team__user", "attachments")
        audit.log(request, "BULK_EXPORT", details=f"{fmt}, {len(ids)} inspections, filters={f}")
        stamp = timezone.localtime().strftime("%Y%m%d_%H%M")
        if fmt == "csv":
            resp = HttpResponse(long_export_csv(insps), content_type="text/csv; charset=utf-8")
            resp["Content-Disposition"] = f'attachment; filename="ITI_inspections_{stamp}.csv"'
        else:
            resp = HttpResponse(long_export_xlsx(insps),
                                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            resp["Content-Disposition"] = f'attachment; filename="ITI_inspections_{stamp}.xlsx"'
        return resp
    page = Paginator(qs, 50).get_page(request.GET.get("page"))
    for i in page:
        i.leader_obj = next((t.user for t in i.team.all() if t.role == TeamMember.LEADER), None)
    qd = request.GET.copy()
    qd.pop("page", None)
    return render(request, "inspections/records.html", {"page": page, "f": f, "qs": qd.urlencode(),
                                                         **_filter_choices()})


# ------------------------------------------------------------------ users
@login_required
@admin_required
def users_list(request):
    q = (request.GET.get("q") or "").strip()
    qs = User.objects.select_related("profile").order_by("-is_active", "first_name", "username")
    if q:
        qs = qs.filter(Q(username__icontains=q) | Q(first_name__icontains=q) | Q(last_name__icontains=q)
                       | Q(profile__designation__icontains=q) | Q(profile__office__icontains=q))
    return render(request, "inspections/users.html", {"users": qs, "q": q})


@login_required
@admin_required
def user_edit(request, pk=None):
    u = get_object_or_404(User, pk=pk) if pk else None
    if request.method == "POST":
        form = UserForm(request.POST, instance=u)
        if form.is_valid():
            if u and u == request.user and form.cleaned_data["role"] != Profile.ADMIN:
                form.add_error("role", "You cannot remove your own administrator role.")
            else:
                saved = form.save()
                audit.log(request, "USER_UPDATED" if u else "USER_CREATED",
                          details=f"{saved.username} role={form.cleaned_data['role']} active={saved.is_active}"
                                  + (" (password reset)" if form.cleaned_data.get("password") else ""))
                messages.success(request, f"User {saved.username} saved.")
                return redirect("users_list")
    else:
        initial = {}
        if u:
            p = getattr(u, "profile", None)
            initial = {"username": u.username, "full_name": u.get_full_name(), "email": u.email,
                       "role": p.role if p else Profile.INSPECTOR, "designation": p.designation if p else "",
                       "office": p.office if p else "", "mobile": p.mobile if p else "", "is_active": u.is_active}
        form = UserForm(initial=initial or {"is_active": True, "role": Profile.INSPECTOR}, instance=u)
    return render(request, "inspections/user_form.html", {"form": form, "obj": u})


# ------------------------------------------------------------------ ITIs
@login_required
@admin_required
def itis_list(request):
    q = (request.GET.get("q") or "").strip()
    qs = ITI.objects.annotate(n=Count("inspections"))
    if q:
        qs = qs.filter(Q(name__icontains=q) | Q(mis_code__icontains=q) | Q(district__icontains=q))
    return render(request, "inspections/itis.html", {"itis": qs, "q": q})


@login_required
@admin_required
def iti_edit(request, pk=None):
    obj = get_object_or_404(ITI, pk=pk) if pk else None
    form = ITIForm(request.POST or None, instance=obj)
    if request.method == "POST" and form.is_valid():
        iti = form.save()
        audit.log(request, "ITI_UPDATED" if obj else "ITI_CREATED", details=f"{iti.mis_code} {iti.name}")
        messages.success(request, f"{iti.name} saved.")
        return redirect("itis_list")
    return render(request, "inspections/simple_form.html", {
        "form": form, "title": "Edit ITI" if obj else "Add ITI", "back": "itis_list"})


def read_table(upload):
    """Read an uploaded .xlsx or .csv into a list of dicts with normalised headers."""
    name = upload.name.lower()
    rows = []
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(upload, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        headers = [str(h or "").strip().lower().replace(" ", "_") for h in next(it)]
        for r in it:
            if r and any(c not in (None, "") for c in r):
                rows.append({headers[i]: ("" if c is None else c) for i, c in enumerate(r) if i < len(headers)})
    elif name.endswith(".csv"):
        text = upload.read().decode("utf-8-sig")
        rd = csv.DictReader(io.StringIO(text))
        for r in rd:
            if any((v or "").strip() for v in r.values()):
                rows.append({(k or "").strip().lower().replace(" ", "_"): (v or "").strip() for k, v in r.items()})
    else:
        raise ValueError("Please upload an .xlsx or .csv file.")
    return rows


ITI_COLUMNS = ["mis_code", "name", "ownership", "district", "block", "address",
               "principal_name", "principal_mobile", "principal_email"]
ASSIGN_COLUMNS = ["mis_code", "team_name", "scheduled_date", "deadline", "leader_username", "member_usernames"]


def _template_response(columns, example, fname):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.append(columns)
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.append(example)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 22
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{fname}"'
    return resp


@login_required
@admin_required
def iti_template(request):
    return _template_response(ITI_COLUMNS, ["PR02000123", "Govt. ITI Example", "GOVT", "Mandi", "Sundernagar",
                                            "Village/Town, P.O., Tehsil", "Principal name", "9800000000",
                                            "principal@example.gov.in"], "ITI_master_template.xlsx")


@login_required
@admin_required
def iti_import(request):
    form = ImportForm(request.POST or None, request.FILES or None)
    report = None
    if request.method == "POST" and form.is_valid():
        try:
            rows = read_table(form.cleaned_data["file"])
        except Exception as e:  # noqa: BLE001
            messages.error(request, f"Could not read the file: {e}")
            rows = None
        if rows is not None:
            report = {"created": 0, "updated": 0, "errors": []}
            with transaction.atomic():
                for n, r in enumerate(rows, start=2):
                    code = str(r.get("mis_code", "")).strip()
                    name = str(r.get("name", "")).strip()
                    district = str(r.get("district", "")).strip()
                    if not code or not name or not district:
                        report["errors"].append(f"Row {n}: mis_code, name and district are required.")
                        continue
                    own = str(r.get("ownership", "GOVT")).strip().upper()
                    own = ITI.PRIVATE if own.startswith("P") else ITI.GOVT
                    vals = {"name": name, "district": district, "ownership": own,
                            "block": str(r.get("block", "")).strip(), "address": str(r.get("address", "")).strip(),
                            "principal_name": str(r.get("principal_name", "")).strip(),
                            "principal_mobile": str(r.get("principal_mobile", "")).strip(),
                            "principal_email": str(r.get("principal_email", "")).strip()}
                    _, created = ITI.objects.update_or_create(mis_code=code, defaults=vals)
                    report["created" if created else "updated"] += 1
            audit.log(request, "ITI_IMPORT", details=f"created {report['created']}, updated {report['updated']}, errors {len(report['errors'])}")
    return render(request, "inspections/import.html", {
        "form": form, "report": report, "title": "Import ITI list", "columns": ITI_COLUMNS,
        "template_url": "iti_template", "back": "itis_list",
        "note": "Existing ITIs are matched by MIS code and updated. Ownership: GOVT or PRIVATE."})


# ------------------------------------------------------------------ assignments
@login_required
@admin_required
def assignments_list(request):
    qs, f = filter_inspections(request)
    qs = qs.order_by("scheduled_date", "iti__name")
    for i in qs:
        i.leader_obj = next((t.user for t in i.team.all() if t.role == TeamMember.LEADER), None)
        i.member_objs = [t.user for t in i.team.all() if t.role == TeamMember.MEMBER]
    return render(request, "inspections/assignments.html", {"items": qs, "f": f, **_filter_choices()})


def _set_team(insp, leader, members):
    keep_sections = {t.user_id: t.sections for t in insp.team.all()}
    insp.team.exclude(user__in=[leader, *members]).delete()
    TeamMember.objects.update_or_create(inspection=insp, user=leader,
                                        defaults={"role": TeamMember.LEADER, "sections": []})
    for m in members:
        TeamMember.objects.update_or_create(inspection=insp, user=m,
                                            defaults={"role": TeamMember.MEMBER,
                                                      "sections": keep_sections.get(m.id, [])})


@login_required
@admin_required
def assignment_edit(request, pk=None):
    insp = get_object_or_404(Inspection, pk=pk) if pk else None
    if request.method == "POST":
        form = AssignmentForm(request.POST, instance=insp)
        if form.is_valid():
            with transaction.atomic():
                obj = form.save(commit=False)
                if not insp:
                    obj.created_by = request.user
                obj.save()
                _set_team(obj, form.cleaned_data["leader"], list(form.cleaned_data["members"]))
                seed_prefill(obj, refresh_team=True)
                obj.save()
            audit.log(request, "ASSIGNMENT_UPDATED" if insp else "ASSIGNMENT_CREATED", obj,
                      details=f"team={obj.team_name}; leader={form.cleaned_data['leader'].username}; "
                              f"members={','.join(m.username for m in form.cleaned_data['members'])}; "
                              f"date={obj.scheduled_date}; deadline={obj.deadline}")
            messages.success(request, "Assignment saved.")
            return redirect("assignment_permissions", pk=obj.pk)
    else:
        initial = {}
        if insp:
            leader = insp.leader()
            initial = {"leader": leader, "members": [t.user for t in insp.team.filter(role=TeamMember.MEMBER)]}
        form = AssignmentForm(instance=insp, initial=initial)
    return render(request, "inspections/assignment_form.html", {"form": form, "obj": insp})


@login_required
@admin_required
def assignment_permissions(request, pk):
    insp = get_object_or_404(Inspection.objects.select_related("iti"), pk=pk)
    members = list(insp.team.filter(role=TeamMember.MEMBER).select_related("user", "user__profile"))
    parts = [p for p in parts_list() if not (p["govt_only"] and not insp.iti.is_govt)]
    if request.method == "POST":
        for tm in members:
            if request.POST.get(f"all_{tm.id}"):
                tm.sections = []
            else:
                tm.sections = [p["key"] for p in parts if request.POST.get(f"p_{tm.id}_{p['key']}")]
            tm.save(update_fields=["sections"])
        audit.log(request, "PERMISSIONS_UPDATED", insp,
                  details="; ".join(f"{tm.user.username}: {','.join(tm.sections) or 'all'}" for tm in members))
        messages.success(request, "Section permissions saved.")
        return redirect("assignments_list")
    return render(request, "inspections/assignment_permissions.html",
                  {"insp": insp, "members": members, "parts": parts, "leader": insp.leader()})


@require_POST
@login_required
@admin_required
def assignment_delete(request, pk):
    insp = get_object_or_404(Inspection, pk=pk)
    if insp.status != Inspection.NOT_STARTED:
        messages.error(request, "Only assignments that have not started can be deleted.")
    else:
        audit.log(request, "ASSIGNMENT_DELETED", None, details=f"{insp.iti} {insp.scheduled_date}")
        insp.delete()
        messages.success(request, "Assignment deleted.")
    return redirect("assignments_list")


@login_required
@admin_required
def assignment_template(request):
    return _template_response(ASSIGN_COLUMNS, ["PR02000123", "Team 1 (Mandi)", "2026-11-10", "2026-11-15",
                                               "leader.login", "member1.login, member2.login"],
                              "Inspection_schedule_template.xlsx")


@login_required
@admin_required
def assignment_import(request):
    form = ImportForm(request.POST or None, request.FILES or None)
    report = None
    if request.method == "POST" and form.is_valid():
        try:
            rows = read_table(form.cleaned_data["file"])
        except Exception as e:  # noqa: BLE001
            messages.error(request, f"Could not read the file: {e}")
            rows = None
        if rows is not None:
            report = {"created": 0, "updated": 0, "errors": []}
            users = {u.username.lower(): u for u in User.objects.filter(is_active=True)}
            for n, r in enumerate(rows, start=2):
                try:
                    code = str(r.get("mis_code", "")).strip()
                    iti = ITI.objects.filter(mis_code=code).first()
                    if not iti:
                        raise ValueError(f"no ITI with MIS code {code!r}")
                    sd = parse_date(r.get("scheduled_date"))
                    dl = parse_date(r.get("deadline"))
                    if dl < sd:
                        raise ValueError("deadline is before the scheduled date")
                    leader = users.get(str(r.get("leader_username", "")).strip().lower())
                    if not leader:
                        raise ValueError("leader_username not found")
                    members = []
                    for m in str(r.get("member_usernames", "") or "").replace(";", ",").split(","):
                        m = m.strip().lower()
                        if m:
                            if m not in users:
                                raise ValueError(f"member {m!r} not found")
                            if users[m] != leader:
                                members.append(users[m])
                    team_name = str(r.get("team_name", "")).strip() or f"Team — {leader.get_full_name() or leader.username}"
                    with transaction.atomic():
                        insp = Inspection.objects.filter(iti=iti, scheduled_date=sd).first()
                        created = insp is None
                        if created:
                            insp = Inspection(iti=iti, scheduled_date=sd, created_by=request.user)
                        elif insp.status == Inspection.SUBMITTED:
                            raise ValueError("an inspection on this date is already submitted")
                        insp.deadline = dl
                        insp.team_name = team_name
                        insp.save()
                        _set_team(insp, leader, members)
                        seed_prefill(insp, refresh_team=True)
                        insp.save()
                    report["created" if created else "updated"] += 1
                except Exception as e:  # noqa: BLE001
                    report["errors"].append(f"Row {n}: {e}")
            audit.log(request, "SCHEDULE_IMPORT", details=f"created {report['created']}, updated {report['updated']}, errors {len(report['errors'])}")
    return render(request, "inspections/import.html", {
        "form": form, "report": report, "title": "Import inspection schedule", "columns": ASSIGN_COLUMNS,
        "template_url": "assignment_template", "back": "assignments_list",
        "note": "Dates as YYYY-MM-DD or DD-MM-YYYY. Several members are separated by commas. "
                "A row for the same ITI and scheduled date updates the existing assignment."})


# ------------------------------------------------------------------ settings
@login_required
@admin_required
def portal_settings(request):
    ps = PortalSettings.get()
    units, _, parts = build_index()
    if request.method == "POST":
        form = SettingsForm(request.POST)
        if form.is_valid():
            valid = {u.key for u in units}
            ps.max_upload_mb = form.cleaned_data["max_upload_mb"]
            ps.allowed_types = ",".join(form.cleaned_data["allowed_types"])
            ps.mandatory_units = [k for k in request.POST.getlist("mandatory") if k in valid]
            ps.save()
            audit.log(request, "SETTINGS_UPDATED",
                      details=f"max {ps.max_upload_mb} MB; types {ps.allowed_types}; mandatory {len(ps.mandatory_units)} items")
            messages.success(request, "Settings saved.")
            return redirect("portal_settings")
    else:
        form = SettingsForm(initial={"max_upload_mb": ps.max_upload_mb, "allowed_types": ps.allowed_ext_list()})
    mand = set(ps.mandatory_units)
    grouped = []
    for p in parts:
        grouped.append({"part": p, "units": [{"u": u, "on": u.key in mand} for u in units if u.part_key == p["key"]]})
    return render(request, "inspections/settings.html", {"form": form, "grouped": grouped, "n_mand": len(mand)})


@login_required
@admin_required
def audit_view(request):
    qs = AuditLog.objects.select_related("user", "inspection__iti")
    f = {k: (request.GET.get(k) or "").strip() for k in ("user", "action", "inspection")}
    if f["user"]:
        qs = qs.filter(Q(user__username__icontains=f["user"]) | Q(user__first_name__icontains=f["user"]))
    if f["action"]:
        qs = qs.filter(action=f["action"])
    if f["inspection"].isdigit():
        qs = qs.filter(inspection_id=int(f["inspection"]))
    page = Paginator(qs, 100).get_page(request.GET.get("page"))
    actions = AuditLog.objects.values_list("action", flat=True).distinct().order_by("action")
    qd = request.GET.copy()
    qd.pop("page", None)
    return render(request, "inspections/audit.html", {"page": page, "f": f, "actions": actions, "qs": qd.urlencode()})
